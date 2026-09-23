"""Build a MEL plan .docx in the Log Frame export format (landscape)."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from app.modules.assess.services.mel_logframe import build_logframe_plan
from app.modules.assess.services.mel_playbooks import parse_playbook_links

_NAVY = RGBColor(0x0D, 0x2C, 0x4C)
_BLUE = RGBColor(0x1B, 0x75, 0xE0)
_STEEL = RGBColor(0x56, 0x64, 0x6F)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_cell_shading(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:shd"):
            tc_pr.remove(child)
    tc_pr.append(shading)


def _set_cell_text(cell, text: str, *, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    paragraphs = text.split("\n") if text else [""]
    for i, line in enumerate(paragraphs):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)


def _add_hyperlink(paragraph, text: str, url: str, *, size: int = 8) -> None:
    """Append an external hyperlink run to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1B75E0")
    r_pr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size * 2))
    r_pr.append(sz)

    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)

    new_run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _set_cell_playbook_links(
    cell, raw: str, links: list[dict[str, str]] | None = None, *, size: int = 8
) -> None:
    cell.text = ""
    resolved = links if links is not None else parse_playbook_links(raw)
    if not resolved:
        _set_cell_text(cell, raw or "", size=size)
        return
    for i, link in enumerate(resolved):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        _add_hyperlink(p, link["label"], link["url"], size=size)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)


def _add_body_paragraph(doc: Document, text: str, *, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.italic = italic
    p.paragraph_format.space_after = Pt(8)


def _set_paragraph_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def _add_heading(doc: Document, text: str, level: int = 1, *, tight: bool = False, keep_next: bool = False) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = _NAVY
    if tight:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    else:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    if keep_next:
        _set_paragraph_keep_with_next(p)


def _add_table_section_title(table, title: str, *, cols: int, fill: str = "0D2C4C") -> None:
    """Insert a full-width title row at the top of a table (stays with headers)."""
    # python-docx can't easily insert before row 0; callers build title as row 0.
    cell = table.rows[0].cells[0]
    if cols > 1:
        cell.merge(table.rows[0].cells[cols - 1])
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = _WHITE
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    _set_cell_shading(cell, fill)


def _set_landscape(doc: Document, *, margin_cm: float = 1.4) -> None:
    """Force landscape A4 with explicit page size (avoids width/height mismatch)."""
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    m = Cm(margin_cm)
    section.left_margin = m
    section.right_margin = m
    top = Cm(1.2) if margin_cm > 0 else Cm(0)
    section.top_margin = top
    section.bottom_margin = top


def _load_title_font(size: int):
    from PIL import ImageFont

    path = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    try:
        return ImageFont.truetype(path, size)
    except OSError:
        return ImageFont.load_default()


def _build_title_page_png(
    *,
    heading: str,
    project_name: str,
    plan_name: str = "",
    intervention_name: str = "",
    plans_included: int | None = None,
) -> BytesIO:
    """Raster title page — reliable full-bleed in Word and Pages."""
    from PIL import Image, ImageDraw

    # A4 landscape @ 150 dpi
    w, h = 1754, 1240
    img = Image.new("RGB", (w, h), (0x1B, 0x75, 0xE0))
    draw = ImageDraw.Draw(img)

    font_eyebrow = _load_title_font(28)
    font_title = _load_title_font(64)
    font_label = _load_title_font(22)
    font_value = _load_title_font(36)

    cx, cy = w // 2, h // 2

    def _centered(text: str, y: int, font, fill: tuple[int, int, int]) -> None:
        bbox = draw.textbbox((0, 0), text, font=font)
        tw = bbox[2] - bbox[0]
        draw.text((cx - tw / 2, y), text, font=font, fill=fill)

    _centered("ASSESS  ·  MEL PLAN", cy - 150, font_eyebrow, (0xB8, 0xD8, 0xF8))
    _centered(heading, cy - 90, font_title, (255, 255, 255))
    draw.line([(cx - 140, cy + 10), (cx + 140, cy + 10)], fill=(0x7D, 0xC3, 0xFF), width=3)

    meta: list[tuple[str, str]] = [("PROJECT", project_name or "—")]
    if plans_included is not None:
        meta.append(("PLANS INCLUDED", str(plans_included)))
    else:
        if plan_name:
            meta.append(("PLAN", plan_name))
        if intervention_name:
            meta.append(("INTERVENTION", intervention_name))

    y = cy + 50
    for label, value in meta:
        _centered(label, y, font_label, (0xB8, 0xD8, 0xF8))
        y += 32
        _centered(value, y, font_value, (255, 255, 255))
        y += 56

    buf = BytesIO()
    img.save(buf, format="PNG", optimize=True)
    buf.seek(0)
    return buf


def _add_title_page(
    doc: Document,
    *,
    project_name: str,
    plan_name: str = "",
    intervention_name: str = "",
    plans_included: int | None = None,
    title: str | None = None,
) -> None:
    """Blue title page image; body content follows on the next page naturally."""
    heading = title or (
        "Project level MEL plan"
        if plans_included is not None
        else "Intervention MEL plans"
    )

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)

    usable_w = section.page_width - section.left_margin - section.right_margin
    usable_h = section.page_height - section.top_margin - section.bottom_margin

    png = _build_title_page_png(
        heading=heading,
        project_name=project_name,
        plan_name=plan_name,
        intervention_name=intervention_name,
        plans_included=plans_included,
    )
    # Fill the content area. No explicit page break — the image alone fills page 1
    # and the next flowable starts on page 2 (an extra break creates a blank page).
    doc.add_picture(png, width=usable_w, height=usable_h - Pt(2))
    pic_par = doc.paragraphs[-1]
    pic_par.paragraph_format.space_before = Pt(0)
    pic_par.paragraph_format.space_after = Pt(0)


def _add_note_callout(doc: Document, note: str) -> None:
    note = (note or "").strip()
    if not note:
        return
    body = re.sub(r"^Please note:\s*", "", note, flags=re.I)
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    _set_cell_shading(cell, "FFF8EF")
    cell.text = ""
    label_p = cell.paragraphs[0]
    lr = label_p.add_run("NOTE")
    lr.bold = True
    lr.font.name = "Calibri"
    lr.font.size = Pt(9)
    lr.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
    label_p.paragraph_format.space_after = Pt(2)
    body_p = cell.add_paragraph()
    br = body_p.add_run(body)
    br.font.name = "Calibri"
    br.font.size = Pt(10)
    br.font.color.rgb = RGBColor(0x5C, 0x46, 0x30)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)
    body_p.paragraph_format.space_after = Pt(0)


def _append_logframe_table(doc: Document, frame: dict[str, Any]) -> None:
    """Append Log frame section title + outcomes table (title stays with headers)."""
    has_type = bool(frame.get("has_outcome_type"))
    headers = (
        ["Outcome type", "Results Chain", "Indicators", "Assumptions"]
        if has_type
        else ["Results Chain", "Indicators", "Assumptions"]
    )
    table = doc.add_table(rows=2, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    _add_table_section_title(table, "Log frame", cols=len(headers), fill="0D2C4C")
    for i, h in enumerate(headers):
        _set_cell_text(table.rows[1].cells[i], h, bold=True, size=9)
        _set_cell_shading(table.rows[1].cells[i], "1B75E0")
        for p in table.rows[1].cells[i].paragraphs:
            for run in p.runs:
                run.font.color.rgb = _WHITE

    for row in frame.get("outcomes") or []:
        cells = table.add_row().cells
        kind = (row.get("kind") or "outcome").lower()
        label = (row.get("label") or "").strip()
        title = (row.get("title") or "").strip()
        chain = f"{label}\n\n{title}".strip() if label else title
        inds = row.get("indicators") or []
        ind_text = "\n".join(
            f"{i}. {name}" if len(inds) > 1 else name for i, name in enumerate(inds, 1)
        )
        assumptions = (row.get("assumptions") or "").strip()
        if assumptions:
            parts = re.split(r"\n(?=\s*\d+[\.\)]\s)", assumptions)
            if len(parts) > 1:
                assumptions = "\n".join(
                    re.sub(r"^\s*\d+[\.\)]\s*", f"{n}. ", p.strip(), count=1)
                    for n, p in enumerate(parts, 1)
                )
        if has_type:
            values = [
                (row.get("outcome_type") or "").strip(),
                chain,
                ind_text,
                assumptions,
            ]
        else:
            values = [chain, ind_text, assumptions]
        for i, val in enumerate(values):
            chain_col = 1 if has_type else 0
            is_cta = bool(row.get("cta")) or kind == "goal"
            if is_cta and i == chain_col:
                cell = cells[i]
                cell.text = ""
                lines = (val or "").split("\n")
                for li, line in enumerate(lines):
                    p = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                    run = p.add_run(line)
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
                    is_prompt = line.strip().startswith("[")
                    run.bold = (not is_prompt) and bool(line.strip())
                    run.italic = is_prompt
                    if is_prompt and kind == "goal":
                        run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
                    elif is_prompt and kind == "output":
                        run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(0)
            else:
                _set_cell_text(cells[i], val, bold=False, size=9)
            if kind == "goal":
                _set_cell_shading(cells[i], "FFF6E8")
            elif kind == "output":
                _set_cell_shading(cells[i], "E8F7F2")


def _append_methodology_table(doc: Document, frame: dict[str, Any]) -> None:
    """Append methodology section title + indicator table."""
    m_headers = [
        "Indicator",
        "Type of monitoring",
        "Data",
        "Method/Instrumentation",
        "Frequency of monitoring",
        "Playbooks",
    ]
    m_title = frame.get("methodology_heading") or "Data and Methodology for Indicators"
    mtable = doc.add_table(rows=2, cols=len(m_headers))
    mtable.style = "Table Grid"
    _add_table_section_title(mtable, m_title, cols=len(m_headers), fill="0D2C4C")
    for i, h in enumerate(m_headers):
        _set_cell_text(mtable.rows[1].cells[i], h, bold=True, size=8)
        _set_cell_shading(mtable.rows[1].cells[i], "1B75E0")
        for p in mtable.rows[1].cells[i].paragraphs:
            for run in p.runs:
                run.font.color.rgb = _WHITE

    for row in frame.get("methodology") or []:
        cells = mtable.add_row().cells
        values = [
            row.get("indicator") or "",
            row.get("monitoring_type") or "",
            row.get("data") or "",
            row.get("method") or "",
            row.get("frequency") or "",
        ]
        for i, val in enumerate(values):
            _set_cell_text(cells[i], val, bold=(i == 0), size=8)
        _set_cell_playbook_links(
            cells[5],
            row.get("playbooks") or "",
            row.get("playbook_links"),
            size=8,
        )


def build_mel_plan_docx(
    *,
    project_name: str,
    plan_name: str,
    intervention_name: str,
    outcomes: list[dict[str, Any]],
    intervention_slug: str | None = None,
) -> bytes:
    """Render the log-frame Word export (landscape A4).

    ``outcomes`` should be the selected catalog outcomes (titles used to filter
    the log-frame template).
    """
    slug = intervention_slug or ""
    try:
        frame = build_logframe_plan(
            intervention_slug=slug or intervention_name,
            selected_outcomes=outcomes,
        )
    except KeyError:
        doc = Document()
        _set_landscape(doc)
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        _add_title_page(
            doc,
            project_name=project_name,
            plan_name=plan_name,
            intervention_name=intervention_name,
        )
        for outcome in outcomes:
            doc.add_heading(outcome.get("title") or "Outcome", level=1)
            if outcome.get("assumptions"):
                doc.add_heading("Assumptions", level=2)
                doc.add_paragraph(outcome["assumptions"])
            for ind in outcome.get("indicators") or []:
                title = ind.get("title") if isinstance(ind, dict) else str(ind)
                if title:
                    doc.add_paragraph(title, style="List Bullet")
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

    doc = Document()
    _set_landscape(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    _add_title_page(
        doc,
        project_name=project_name,
        plan_name=plan_name,
        intervention_name=intervention_name,
    )

    _add_heading(doc, "Overview", level=1)
    for para in frame.get("intro") or []:
        _add_body_paragraph(doc, para)

    _append_logframe_table(doc, frame)
    _add_note_callout(doc, (frame.get("note") or "").strip())
    _append_methodology_table(doc, frame)

    doc.add_paragraph()
    _append_outro_sections(doc, frame.get("outro") or [])

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_mel_project_docx(
    *,
    project_name: str,
    plans: list[dict[str, Any]],
) -> bytes:
    """Build one landscape Word doc for all MEL plans in a project.

    Shared overview / frequency / field-setup prose appears once. Each plan
    contributes its log frame + methodology (+ note).
    """
    from app.modules.assess.services.mel_logframe import resolve_selectable_outcomes

    frames: list[tuple[dict[str, Any], dict[str, Any]]] = []
    for plan in plans:
        slug = plan.get("intervention_slug") or ""
        outcome_ids = (plan.get("plan_json") or {}).get("outcome_ids") or []
        try:
            outcomes = resolve_selectable_outcomes(slug, outcome_ids)
            frame = build_logframe_plan(
                intervention_slug=slug,
                selected_outcomes=outcomes,
                selected_outcome_ids=outcome_ids,
            )
        except KeyError:
            continue
        frames.append((plan, frame))

    if not frames:
        raise ValueError("No exportable MEL plans with log-frame templates in this project")

    doc = Document()
    _set_landscape(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    _add_title_page(
        doc,
        project_name=project_name,
        plans_included=len(frames),
    )

    shared_intro = frames[0][1].get("intro") or []
    _add_heading(doc, "Overview", level=1)
    for para in shared_intro:
        _add_body_paragraph(doc, para)

    for plan, frame in frames:
        name = frame.get("name") or plan.get("name") or plan.get("intervention_slug") or ""
        plan_label = plan.get("name") or name
        _add_heading(doc, str(name), level=1, tight=True)
        meta = doc.add_paragraph()
        mr = meta.add_run(f"Plan: {plan_label}")
        mr.font.name = "Calibri"
        mr.font.size = Pt(10)
        mr.font.color.rgb = _STEEL
        meta.paragraph_format.space_after = Pt(6)

        _append_logframe_table(doc, frame)
        _add_note_callout(doc, (frame.get("note") or "").strip())
        _append_methodology_table(doc, frame)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)

    _append_outro_sections(doc, frames[0][1].get("outro") or [])

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


_SECTION_HEADS = {
    "frequency",
    "sampling",
    "field setup for data collections",
    "what do you do with the mel plan?",
    "what do you do with the mel plan",
}


def _norm_head(text: str) -> str:
    return re.sub(r"\.$", "", (text or "").strip().lower())


def _is_section_head(text: str) -> bool:
    return _norm_head(text) in _SECTION_HEADS


def _is_sampling_bullet(text: str) -> bool:
    return bool(
        text
        and (
            text.startswith("For watershed management programmes")
            or text.startswith("Within each watershed")
            or text.startswith("Zones that are distinct")
        )
    )


def _is_field_step(text: str) -> bool:
    return bool(
        text
        and (
            text.startswith("Select control assets:")
            or text.startswith("Asset allocation:")
            or text.startswith("Deploy Instruments:")
            or text.startswith("Train community")
        )
    )


def _append_outro_sections(doc: Document, outro: list[str]) -> None:
    """Render Frequency + Field setup with numbered lists and clear headings."""
    items = list(outro or [])
    freq_idx = next((i for i, p in enumerate(items) if _norm_head(p) == "frequency"), -1)
    after = next(
        (
            i
            for i, p in enumerate(items)
            if i > freq_idx
            and (_norm_head(p).startswith("what do you do") or _norm_head(p) == "sampling")
        ),
        -1,
    )
    if freq_idx >= 0 and after > freq_idx:
        frequency = items[freq_idx:after]
        field = items[after:]
    elif freq_idx == 0:
        frequency = items[:2]
        field = items[2:]
    else:
        frequency = []
        field = items

    if frequency:
        _add_heading(doc, "Frequency", level=1)
        for para in frequency:
            if _is_section_head(para) and _norm_head(para) == "frequency":
                continue
            _add_body_paragraph(doc, para)

    _add_heading(doc, "Field setup", level=1)
    i = 0
    while i < len(field):
        para = field[i]
        if _is_section_head(para):
            _add_heading(doc, para.rstrip("."), level=2)
            i += 1
            continue
        if _is_sampling_bullet(para):
            while i < len(field) and _is_sampling_bullet(field[i]):
                p = doc.add_paragraph(field[i], style="List Number")
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_after = Pt(4)
                i += 1
            continue
        if _is_field_step(para):
            while i < len(field) and _is_field_step(field[i]):
                step = field[i]
                colon = step.find(":")
                step_title = step[: colon + 1] if colon >= 0 else step
                body = step[colon + 1 :] if colon >= 0 else ""
                p = doc.add_paragraph(style="List Number")
                title_run = p.add_run(step_title)
                title_run.bold = True
                title_run.font.name = "Calibri"
                title_run.font.size = Pt(11)
                if body:
                    body_run = p.add_run(body)
                    body_run.bold = False
                    body_run.font.name = "Calibri"
                    body_run.font.size = Pt(11)
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_after = Pt(6)
                i += 1
            continue
        _add_body_paragraph(doc, para)
        i += 1
